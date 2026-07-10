# -*- coding: utf-8 -*-
"""AeroForge Final Report — human tone (aligned with the corrected Business Plan voice)."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = "/home/user/my-profile-5/business-simulation"
CH = f"{HERE}/charts"

INK = RGBColor(0x0B, 0x0B, 0x0B)
SEC = RGBColor(0x52, 0x51, 0x4E)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(10.5)
for h in ("Heading 1", "Heading 2"):
    doc.styles[h].font.name = "Calibri"
    doc.styles[h].font.color.rgb = INK
doc.styles["Heading 1"].font.size = Pt(15)
doc.styles["Heading 2"].font.size = Pt(12)

def P(text, bold=False, italic=False, size=10.5, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.bold = bold; r.font.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def H1(t): doc.add_heading(t, level=1)
def H2(t): doc.add_heading(t, level=2)

def table(headers, rows, highlight_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 6"
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        r.font.bold = True; r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9.5)
            if highlight_col is not None and j == highlight_col:
                r.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def fig(path, width, caption):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(caption, italic=True, size=9, color=SEC, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

# ================= Title
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AeroForge — Final Report"); r.font.size = Pt(24); r.font.bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Business Simulation Capstone — Marketplace® Conscious Capitalism (Bikes) · July 2026")
r.font.size = Pt(11); r.font.color.rgb = SEC
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Executive team: Aaron Rainier (President) · Luna Carballo (VP Marketing) · Justine Pellier "
              "(VP Business Analytics) · Patrick Yang (VP Accounting & Finance) · Sarah Baraket (VP Sales) · "
              "Teysnim Abichou (VP Human Resources)")
r.font.size = Pt(10); r.font.color.rgb = SEC
P("", space_after=4)

# ================= 1
H1("1. Mission Statement and Vision")
P("AeroForge exists to put elite carbon performance within everyone's reach. We forge custom-fitted, "
  "3D-printed carbon fiber bikes that empower riders to conquer any trail and win any race — while "
  "creating value for our employees, our communities, and the planet we ride through.", italic=True)
P("We wrote this mission on the first afternoon, before we had sold a single bike, and if we are honest it "
  "stayed mostly words for the first two quarters. From quarter 3 it started to mean something concrete. It "
  "kept us out of a price war when we were last in market share and cutting prices would have been the easy "
  "move. It made us keep spending on our workers and on the environment during quarters when we were losing "
  "money. And when the industry news broke that standard carbon frames were breaking and injuring mountain "
  "riders, the biggest decision of our game — spending $1.02M on enriched carbon fiber — was actually not a "
  "hard one to make. Our mission says we sell bikes that let people conquer any trail; that has to mean bikes "
  "that are safe to ride. By Q6, our entire range was built on enriched carbon.")

# ================= 2
H1("2. Business Strategy Summary")
P("We chose Mountain as our primary segment and Speed as our secondary one, because they were the two most "
  "profitable segments and their needs overlap enough to share technology. In Q4 the team added a third, "
  "opportunistic brand for the Recreation segment (AeroForgeComfy) — outside our core strategy, but cheap to "
  "do and, in the end, a good call. Our positioning never changed: premium bikes, sold on quality and "
  "integrity, never on price. Distribution followed the demand data: New York first, then Amsterdam in Q4 "
  "(the largest market in the game when we committed to it), then Bangalore in Q5. All of this was funded by "
  "$1.5M of founding capital, $1.0M of team tranches, and the $2.5M venture round we obtained on our business "
  "plan — money we spent exactly the way we told the investors we would: R&D, production capacity, stores.")
P("Looking back, our six quarters split into five phases. Quarters 1–2: set up and test the market — this is "
  "where we failed, by designing on intuition. Quarter 3: repair — we rebuilt the Bolt from the market data "
  "and it became the #1-rated Speed bike within one quarter. Quarter 4: invest — the venture money went into "
  "R&D features, more capacity, the Amsterdam store, and a compensation catch-up for our people. Quarter 5: "
  "expand — we doubled the sales force to 13 across two cities, opened Bangalore, shipped our own tire "
  "technologies, and placed the enriched-carbon bet. Quarter 6: harvest — the whole range relaunched on "
  "enriched carbon, and the company posted the first profit of its history: +$708k on $3.0M of revenue.")

# ================= 3
H1("3. Performance Review: Financial, Market, Operational, HR and CSR")

H2("3.1 Financial performance")
P("Our revenue grew ×13 between the test market and Q6: 226k, 373k, 740k, 1,440k, then 3,018k. The net income "
  "line tells the story of the whole company: three small losses while we repaired the product (−257k, −151k, "
  "−106k), one big planned loss when we deployed the venture money (−849k in Q4), an even bigger one in Q5 "
  "when we bought the enriched carbon technology (−1,246k) — and then, in the final quarter, our first "
  "profit: +708k, at a 62.5% gross margin, up from 53% in Q2. Through all of it we never borrowed a dollar: "
  "no bank loans, no emergency loans, and a Financial Risk score of 1.000 — the class maximum — every single "
  "quarter. We finished with $2.06M in cash.")
P("So why did we still finish last on cumulative financial performance (14.4 against a class average of "
  "40.0)? Because the index rewards accumulated profit, and ours arrived one quarter before the game ended. "
  "Our Wealth score says the same thing: −$1.9M of retained earnings against $5.0M of paid-in capital. That "
  "is the bill for the quarter we failed, and for a strategy that deliberately spent the middle of the game "
  "preparing for the end of it.")
fig(f"{CH}/chart_finance.png", 5.6, "Figure 1 — Revenue and operating cash flow by quarter ($ thousands). "
    "The Q4–Q5 troughs are the planned deployment of the $2.5M venture round.")

H2("3.2 Market performance")
P("This is where our early mistakes cost us the most. Our market share went from 7.0% in Q3 to 8.9% in Q6, "
  "and our cumulative Market Performance finished last (0.071 against an average of 0.159). The causes are "
  "well documented and there are exactly two. First, for two quarters our Mountain bike sat below the "
  "Customer Union's minimum rating, so the biggest source of demand in the game simply did not recommend us. "
  "Second, once we fixed the products, we could never build enough of them: we lost sales to stock-outs in "
  "every quarter from Q3 onward — 42, 88, 79, and then 1,083 units in Q6 alone.")
P("The frustrating part is that the product ended the game excellent. Brand judgments rose from 53, 62 and 65 "
  "in the early quarters to 87 (Speed), 83 (Mountain, top-2) and 78 (Recreation, top-2) in Q6. Our "
  "advertising, rebuilt in Q3–Q4, scored 80 and above. Our prices rated at or near 100 throughout. In Q6, "
  "customers asked us for 3,318 bikes — 84% more than we ourselves had dared to project — and we could only "
  "deliver 2,235. We started the game with a product nobody wanted enough, and finished it with a product we "
  "could not build enough of.")
fig(f"{CH}/chart_fixes.png", 5.6, "Figure 2 — Diagnose → fix → measure: every metric we corrected moved.")

H2("3.3 Operations")
P("Production capacity scaled from 6 to 32 units per day across the game (6 → 11 → 16 → 32), with two "
  "3D-printer purchases totalling $960k. Worker productivity climbed from 68.6% to 82.4% and unit labour cost "
  "fell from $108 to $84. Quality investments — statistical process control, operator training, supplier "
  "programs — steadily improved reliability. Our recurring operational mistake was simpler and more human: we "
  "kept underestimating our own demand. Every time we fixed something in marketing, more customers showed up "
  "than we had planned for, four quarters in a row, and roughly $1.4M of Q6 revenue was left on the table. We "
  "kept treating capacity as a cost to be minimised when it had quietly become our real bottleneck.")

H2("3.4 Human resources")
P("We learned this one the hard way. In quarters 2–3 we paid our people below the market. On paper it looked "
  "like savings; in reality it was one of the most expensive decisions of our game — worker satisfaction sat "
  "at 71.9%, a third of our workforce left every quarter, and productivity fell while our competitors' rose. "
  "From Q4 we moved every package to the top of the market: $13.2k base for workers with expanded health "
  "cover, three weeks of vacation and a 5% pension, with equivalent upgrades for supervisors and the sales "
  "force. People stayed (turnover fell from 32% to 11%), worked better (+14 productivity points) and were "
  "visibly happier (satisfaction 88%). Our cumulative HR score finished at 0.790 against a 0.801 average — "
  "those two cheap early quarters are the whole difference.")

H2("3.5 CSR / Conscious Capitalism")
P("By the end of the game we were running sixteen improvement programs: chemical collection and containment, "
  "a clean-room air-filtration facility, protective equipment, air-filtering plants, a carbon-recycling "
  "consortium, cross-training, quality programs with our suppliers, a fitness centre, daycare, school grants "
  "and community bike trails. We want to be honest about where the clean room came from: not from a metric, "
  "but from two memos written by our own production workers telling us they were breathing carbon dust and "
  "epoxy vapour. Our people told us before any number did, and we acted the same quarter. AeroForge was never "
  "cited for toxic leaks — the industry newsletter repeatedly named four of our five competitors. Ethics also "
  "cost us something real twice: we withdrew advertising claims both times the truth-in-advertising check "
  "flagged them, and we declined to buy a licence for a technology we had already developed ourselves rather "
  "than waste shareholder money. Our cumulative Reputation score, 0.679 against a 0.690 average, carries the "
  "trace of starting all of this only in Q4.")

# ================= 4
H1("4. Main Results vs. Other Teams")
P("The tables below compare us with the other five teams, without cosmetics.")
P("Final standings — cumulative Total Performance by company (the official ranking):")
table(
    ["Rank", "Company", "Cumulative Total Perf.", "Q6 Total Perf. (final quarter)"],
    [
        ("1", "WBB", "54.121", "357.724"),
        ("2", "NatuRide", "22.149", "101.200"),
        ("3", "Spe3d", "8.662", "31.720"),
        ("4", "Carbon-Ride", "2.918", "4.534"),
        ("5", "The Bike Yard", "0.968", "3.318"),
        ("6", "AeroForge", "0.703", "3.494"),
    ],
)
P("There are two honest ways to read this table. The first: we finished sixth out of six. The second: the gap "
  "to fifth place is 0.27 points — fourth place sits at 2.9 and first at 54 — and on the final quarter taken "
  "alone we were no longer last. The game stopped at the moment our curve was starting to cross the others.")
P("Final cumulative balanced scorecard, AeroForge against the class:")
table(
    ["Cumulative metric", "Class min", "Class avg", "Class max", "AeroForge"],
    [
        ("Total Performance", "0.703", "14.920", "54.121", "0.703 — last"),
        ("Financial Performance", "14.380", "39.963", "74.217", "14.380 — last"),
        ("Market Performance", "0.071", "0.159", "0.271", "0.071 — last"),
        ("Marketing Effectiveness", "0.619", "0.743", "0.799", "0.745 — above avg"),
        ("Investment in Future", "4.228", "5.223", "7.874", "7.874 — 1st in class"),
        ("Wealth", "0.620", "1.053", "1.460", "0.620 — last"),
        ("HR Management", "0.766", "0.801", "0.854", "0.790"),
        ("Asset Management", "0.487", "0.771", "1.185", "0.487 — last"),
        ("Manufacturing Productivity", "0.636", "0.701", "0.738", "0.720 — above avg"),
        ("Financial Risk", "0.914", "0.986", "1.000", "1.000 — 1st (tied)"),
        ("Reputation", "0.635", "0.690", "0.752", "0.679"),
    ],
    highlight_col=4,
)
P("Quarterly Total Performance — the trajectory behind the cumulative number:")
table(
    ["", "Q2", "Q3", "Q4", "Q5", "Q6"],
    [
        ("AeroForge", "0.013", "0.094", "0.266", "0.957", "3.494"),
        ("Class average", "0.229", "0.943", "4.190", "n/a", "83.7"),
    ],
)
fig(f"{CH}/chart_totalperf.png", 5.4, "Figure 3 — AeroForge quarterly Total Performance: ×269 in four "
    "quarters, from 6% of the class average in Q2 to a level that would have been mid-pack two quarters earlier.")
P("Read together, the tables explain our ranking honestly. The cumulative index is driven by accumulated "
  "financial results, and starting 17 times below the class average in Q2 is not something four quarters can "
  "repair, however steep the recovery. What the same tables also show is that the two metrics a team can win "
  "from any starting position — Investment in Future and Financial Risk — we won, finishing first in class on "
  "both.")

# ================= 5
H1("5. Lessons Learnt")
P("Lesson 1 — Under uncertainty, intuition is a hypothesis, not a decision. Written down, our three Q2 "
  "mistakes look almost silly: a race-bike advertisement full of mountain images and claims (it scored 35/100 "
  "— it actually appealed more to Mountain customers than to the Speed customers it was made for); a comfort "
  "handlebar and comfort seat on a mountain bike, because comfort sounded nice (the bike scored 62, below the "
  "recommendation threshold, so the Customer Union simply never recommended us); and the choice to skip "
  "roughly $45k of market studies to save money. That last one was the real mistake — the data that would "
  "have corrected the first two was on sale the whole time. It cost us two quarters and, through the "
  "cumulative scorecard, the game.", space_after=8)
P("Lesson 2 — The business disciplines are one system. A 35-point ad starves a factory. Below-market pay cuts "
  "productivity, which caps output, which turns your marketing wins into stock-outs. We never once lost money "
  "'in finance': every loss we posted was born upstream, in a marketing or HR decision, and the accounts just "
  "recorded it.", space_after=8)
P("Lesson 3 — Learn how the scorecard is computed before playing, not after. Some metrics can be won from any "
  "position, whatever your sales volume — Investment in Future, Financial Risk — and we did win them, but we "
  "only started managing them deliberately from Q4. The cumulative Total Performance metric, on the other "
  "hand, never forgets a bad quarter. Protecting the floor matters more than raising the ceiling.",
  space_after=8)
P("Lesson 4 — When product quality jumps, demand does not follow politely. It jumps too. We under-projected "
  "demand in all four quarters after our repairs, ending with 3,318 units of Q6 demand against our own "
  "projection of 1,800, because we kept extrapolating in straight lines. Capacity should be built ahead of "
  "proven product quality, not behind last quarter's sales.", space_after=8)
P("Lesson 5 — People tell you before the metrics do. The worker memos about carbon dust arrived before any "
  "visible drop in morale; acting on them came before the gains in productivity and satisfaction. The "
  "in-basket turned out to be our cheapest and fastest market research.", space_after=8)
P("The real turning point was not a single decision but a change in how we worked. From Q3 onward, nobody in "
  "our meetings could defend a choice without a number behind it — a purchased study, a revealed competitor "
  "matrix, an industry table. Every advertisement went through the truth-in-advertising check before "
  "submission. Every quarter closed the same way: Final Check, then the pro forma recomputed last. Each VP "
  "owned one function and the President arbitrated the trade-offs. That routine is what turned a failing "
  "start-up into a company that improved every metric it touched, four quarters in a row.")

# ================= 6
H1("6. Preparing the Firm to Compete in the Future")
P("We know 'we prepared the future' can sound like an excuse from a team that finished last, so we will let "
  "the game's own measurement speak first: AeroForge finished first in class on cumulative Investment in "
  "Future. In concrete terms, a hypothetical Q7 would have opened with the industry's only fully "
  "enriched-carbon range, rated 87/83/78 one quarter after competitors began converting theirs; our own tire "
  "technologies on the road and a further one — hybrid comfort tires for the Recreation line — already paid "
  "for and arriving; a modern 32-units-per-day factory run by a stable, well-paid workforce (82% productivity, "
  "11% turnover); three stores covering the three largest markets with 21 trained sales and service people; "
  "$2.06M in cash, zero debt and untouched borrowing capacity; and sixteen running programs protecting "
  "quality, the environment and our community. Proven Q6 demand of 3,318 units against 2,235 produced means "
  "the first Q7 decision was already obvious: one more printer and a third shift.")

# ================= 7
H1("7. Top Recommendations If We Were to Start Over")
P("1. Spend your first dollar on information. Buy every market study from the test market onward, and design "
  "brands, ads and prices from the data — never from instinct. One failed quarter costs fifty times the price "
  "of the studies.", space_after=6)
P("2. Benchmark everything before you ship it: an ad against the truth-check, a salary against the industry "
  "tables, a price against what customers say they will pay. The market has usually already published the "
  "right answer.", space_after=6)
P("3. Build capacity ahead of product quality, not behind sales. Once your ratings are top-tier, demand "
  "arrives faster than a spreadsheet expects; four quarters of stock-outs handed our best marketing to our "
  "competitors.", space_after=6)
P("4. Pay people at the top of the market from quarter one. The morale-to-productivity loop compounds slowly, "
  "and started late it never fully catches up.", space_after=6)
P("5. Read the scorecard's formulas on day one, and protect the cumulative floor: one collapsed quarter "
  "follows you to the end, because the index never forgets.", space_after=6)
P("6. Keep the financial discipline we kept: no emergency borrowing, investments financed by capital actually "
  "raised on a credible plan. Financial Risk was a perfect score available to every team; most dropped it.",
  space_after=6)
P("We lost this race, and we know exactly why — which is, we believe, the outcome this simulation is designed "
  "to produce. Give the machine we built over the last four quarters a seventh one, and we would have been "
  "the most dangerous company in the industry.", italic=True)

doc.save(f"{REPO}/AeroForge_Final_Report.docx")
print("saved", f"{REPO}/AeroForge_Final_Report.docx")
